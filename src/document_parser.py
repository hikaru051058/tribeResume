"""Document parsing utilities for resume inputs.

The parser normalizes supported resume formats into text plus lightweight
metadata. It does not perform OCR or claim full layout understanding.
"""

from __future__ import annotations

from pathlib import Path
from typing import Any


SUPPORTED_TYPES = {
    ".txt": "txt",
    ".md": "markdown",
    ".markdown": "markdown",
    ".docx": "docx",
    ".pdf": "pdf",
}


def _word_count(text: str) -> int:
    return len(text.split())


def _base_result(
    path: str | Path,
    file_type: str,
    text: str = "",
    metadata: dict[str, Any] | None = None,
    warnings: list[str] | None = None,
) -> dict:
    metadata = metadata or {}
    warnings = warnings or []
    metadata.setdefault("char_count", len(text))
    metadata.setdefault("word_count", _word_count(text))
    if metadata["char_count"] < 200 or metadata["word_count"] < 30:
        warnings.append(
            "Extracted text is very short. The file may be incomplete, scanned, or image-based."
        )
    return {
        "source_path": str(path),
        "file_type": file_type,
        "text": text,
        "metadata": metadata,
        "warnings": warnings,
    }


def detect_file_type(path: str) -> str:
    """Detect supported file type from file extension."""

    file_path = Path(path)
    suffix = file_path.suffix.lower()
    return SUPPORTED_TYPES.get(suffix, "unsupported")


def _read_text_with_fallback(path: str | Path) -> str:
    file_path = Path(path)
    try:
        return file_path.read_text(encoding="utf-8")
    except UnicodeDecodeError:
        return file_path.read_text(encoding="utf-8", errors="replace")


def parse_text_file(path: str) -> dict:
    """Parse a UTF-8 text resume."""

    text = _read_text_with_fallback(path)
    return _base_result(
        path,
        "txt",
        text,
        metadata={"extraction_method": "utf8_text"},
    )


def parse_markdown_file(path: str) -> dict:
    """Parse a Markdown resume as text."""

    text = _read_text_with_fallback(path)
    return _base_result(
        path,
        "markdown",
        text,
        metadata={"extraction_method": "utf8_markdown_text"},
    )


def parse_docx_file(path: str) -> dict:
    """Parse a DOCX resume, including paragraphs and table cell text."""

    from docx import Document

    document = Document(path)
    chunks: list[str] = []

    paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs]
    chunks.extend(text for text in paragraphs if text)

    table_cell_count = 0
    for table in document.tables:
        for row in table.rows:
            row_values = []
            for cell in row.cells:
                table_cell_count += 1
                cell_text = " ".join(
                    paragraph.text.strip()
                    for paragraph in cell.paragraphs
                    if paragraph.text.strip()
                )
                if cell_text:
                    row_values.append(cell_text)
            if row_values:
                chunks.append(" | ".join(row_values))

    text = "\n".join(chunks)
    return _base_result(
        path,
        "docx",
        text,
        metadata={
            "extraction_method": "python-docx",
            "paragraph_count": len([text for text in paragraphs if text]),
            "table_count": len(document.tables),
            "table_cell_count": table_cell_count,
        },
    )


def parse_pdf_file(path: str) -> dict:
    """Parse a text-based PDF resume page by page."""

    from pypdf import PdfReader

    reader = PdfReader(path)
    page_texts = []
    empty_pages = []
    for index, page in enumerate(reader.pages, start=1):
        page_text = page.extract_text() or ""
        page_text = page_text.strip()
        if not page_text:
            empty_pages.append(index)
        page_texts.append(page_text)

    text = "\n\n".join(page_text for page_text in page_texts if page_text)
    warnings = []
    if empty_pages:
        warnings.append(
            "PDF pages with no extracted text: "
            + ", ".join(str(page) for page in empty_pages)
            + ". Scanned/image PDFs are not supported yet."
        )

    return _base_result(
        path,
        "pdf",
        text,
        metadata={
            "extraction_method": "pypdf",
            "page_count": len(reader.pages),
            "empty_text_pages": empty_pages,
        },
        warnings=warnings,
    )


def parse_document(path: str) -> dict:
    """Parse a supported resume document into a consistent dictionary."""

    file_path = Path(path)
    file_type = detect_file_type(str(file_path))

    if not file_path.exists():
        return _base_result(
            file_path,
            file_type,
            metadata={"extraction_method": "none"},
            warnings=[f"File does not exist: {file_path}"],
        )

    if file_type == "txt":
        return parse_text_file(str(file_path))
    if file_type == "markdown":
        return parse_markdown_file(str(file_path))
    if file_type == "docx":
        return parse_docx_file(str(file_path))
    if file_type == "pdf":
        return parse_pdf_file(str(file_path))

    return _base_result(
        file_path,
        "unsupported",
        metadata={"extraction_method": "none"},
        warnings=[
            f"Unsupported file type '{file_path.suffix}'. "
            "Supported formats: .txt, .md, .docx, .pdf."
        ],
    )

